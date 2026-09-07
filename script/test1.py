from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = REPO_ROOT / "test_file" / "demo1.xlsx"
OUTPUT_DIR = REPO_ROOT / "test_file" / "json_data"
OUTPUT_PATH = OUTPUT_DIR / "demo1.json"


def clean_text(value) -> str:
    """Convert a cell value to clean text."""
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", text).strip()


def ocr_image(image) -> str:
    """OCR an embedded Excel image and return only the recognized text."""
    try:
        import pytesseract
        from PIL import Image

        # openpyxl stores the embedded image bytes in _data().
        image_data = image._data()
        with Image.open(__import__("io").BytesIO(image_data)) as pil_image:
            languages = []
            try:
                available = pytesseract.get_languages(config="")
                for lang in ("eng", "jpn", "chi_sim"):
                    if lang in available:
                        languages.append(lang)
            except Exception:
                pass

            lang = "+".join(languages) or "eng"
            return clean_text(pytesseract.image_to_string(pil_image, lang=lang))
    except ImportError:
        return ""
    except Exception as exc:
        print(f"Warning: OCR failed: {exc}")
        return ""


def image_row(image) -> int:
    """Return the 1-based Excel row where the image is anchored."""
    try:
        return image.anchor._from.row + 1
    except AttributeError:
        return 1


def extract_workbook() -> list[dict[str, object]]:
    """Read every sheet in order and return sheet/line/text records."""
    workbook = load_workbook(XLSX_PATH, data_only=True)
    records: list[dict[str, object]] = []

    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        # Map OCR text to the row containing each embedded image.
        image_text_by_row: dict[int, list[str]] = {}
        for image in getattr(worksheet, "_images", []):
            text = ocr_image(image)
            if text:
                image_text_by_row.setdefault(image_row(image), []).append(text)

        max_row = max(worksheet.max_row, max(image_text_by_row.keys(), default=0))

        for row_number in range(1, max_row + 1):
            cell_texts = [
                text
                for cell in worksheet[row_number]
                if (text := clean_text(cell.value))
            ]
            cell_text = " | ".join(cell_texts)
            image_text = " | ".join(image_text_by_row.get(row_number, []))

            parts = [part for part in (cell_text, image_text) if part]
            if not parts:
                continue

            records.append(
                {
                    "sheet": sheet_index,
                    "line": row_number,
                    "text": " | ".join(parts),
                }
            )

    return records


def main() -> None:
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Excel file not found: {XLSX_PATH}")

    # mkdir also creates json_data when it does not exist.
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Always replace the old JSON contents with a fresh extraction.
    records = extract_workbook()
    with OUTPUT_PATH.open("w", encoding="utf-8") as file:
        json.dump(records, file, ensure_ascii=False, indent=2)
        file.write("\n")

    print(f"Extracted {len(records)} lines from {XLSX_PATH}")
    print(f"Output: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
