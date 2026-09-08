from __future__ import annotations

import json
import posixpath
import sys
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / "test_file" / "demo1.xlsx"
OUTPUT_DIR = ROOT / "test_file" / "json_data"
OUTPUT_PATH = OUTPUT_DIR / "demo1.json"

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
DRAWING_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"


def column_letter(column_number: int) -> str:
    """Convert a 1-based Excel column number to A, B, ..., AA, AB, ..."""
    result = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def normalize_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def resolve_target(source_path: str, target: str) -> str:
    """Resolve an OOXML relationship target to a normalized package path."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(
        posixpath.join(posixpath.dirname(source_path), target)
    ).lstrip("./")


def get_sheet_paths(xlsx_path: Path) -> list[tuple[str, str]]:
    """Return (sheet name, worksheet XML path) in workbook order."""
    result = []

    with ZipFile(xlsx_path) as zf:
        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        workbook_rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in workbook_rels}

        sheets = workbook.find(f"{{{NS_MAIN}}}sheets")
        if sheets is None:
            return result

        for sheet in sheets:
            name = sheet.attrib.get("name", "")
            rid = sheet.attrib.get(f"{{{NS_REL}}}id")
            target = rel_map.get(rid)
            if target:
                result.append((name, resolve_target("xl/workbook.xml", target)))

    return result


def get_drawing_texts(xlsx_path: Path) -> dict[str, list[tuple[int, int, str]]]:
    """Read Shape/TextBox text and its starting cell without OCR.

    Only xdr:sp (DrawingML shapes/text boxes) are processed. xdr:pic image
    objects are ignored, so text is not extracted from images.
    """
    result: dict[str, list[tuple[int, int, str]]] = {}

    with ZipFile(xlsx_path) as zf:
        names = set(zf.namelist())
        sheet_paths = get_sheet_paths(xlsx_path)

        for _, sheet_path in sheet_paths:
            rels_path = posixpath.join(
                posixpath.dirname(sheet_path),
                "_rels",
                posixpath.basename(sheet_path) + ".rels",
            )
            if rels_path not in names:
                continue

            sheet_rels = ET.fromstring(zf.read(rels_path))
            drawing_path = None
            for rel in sheet_rels:
                if rel.attrib.get("Type") == DRAWING_REL_TYPE:
                    target = rel.attrib.get("Target")
                    if target:
                        candidate = resolve_target(sheet_path, target)
                        if candidate in names:
                            drawing_path = candidate
                    break

            if not drawing_path:
                continue

            try:
                root = ET.fromstring(zf.read(drawing_path))
            except (KeyError, ET.ParseError):
                continue

            entries = []
            for anchor in root:
                # Only process actual DrawingML shapes/text boxes.
                shape = anchor.find(f"{{{NS_XDR}}}sp")
                if shape is None:
                    continue

                from_node = anchor.find(f"{{{NS_XDR}}}from")
                if from_node is None:
                    continue

                row_node = from_node.find(f"{{{NS_XDR}}}row")
                col_node = from_node.find(f"{{{NS_XDR}}}col")
                if row_node is None or col_node is None:
                    continue

                try:
                    line = int(row_node.text or 0) + 1
                    column = int(col_node.text or 0) + 1
                except ValueError:
                    continue

                tx_body = shape.find(f"{{{NS_XDR}}}txBody")
                if tx_body is None:
                    continue

                texts = [
                    node.text
                    for node in tx_body.iter(f"{{{NS_A}}}t")
                    if node.text
                ]
                text = normalize_text("".join(texts))
                if text:
                    entries.append((line, column, text))

            if entries:
                result[sheet_path] = entries

    return result


def extract_workbook() -> list[dict]:
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Excel file not found: {XLSX_PATH}")

    workbook = load_workbook(XLSX_PATH, data_only=True, read_only=False)
    drawing_texts = get_drawing_texts(XLSX_PATH)
    sheet_paths = dict(get_sheet_paths(XLSX_PATH))

    sheets = []
    for ws in workbook.worksheets:
        context = []

        # Normal worksheet cell text.
        for row in ws.iter_rows():
            for cell in row:
                text = normalize_text(cell.value)
                if text:
                    context.append({
                        "cell": f"{column_letter(cell.column)}{cell.row}",
                        "text": text,
                    })

        # Shape/TextBox text. Use the shape's top-left/start cell.
        sheet_xml = sheet_paths.get(ws.title)
        for line, column, text in drawing_texts.get(sheet_xml, []):
            context.append({
                "cell": f"{column_letter(column)}{line}",
                "text": text,
            })

        def sort_key(item):
            cell = item["cell"]
            column = "".join(c for c in cell if c.isalpha())
            line = int("".join(c for c in cell if c.isdigit()))
            column_number = 0
            for char in column:
                column_number = column_number * 26 + ord(char) - 64
            return line, column_number

        context.sort(key=sort_key)
        sheets.append({
            "name": ws.title,
            "context": context,
        })

    return sheets


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    data = {"sheet": extract_workbook()}
    OUTPUT_PATH.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"Generated: {OUTPUT_PATH}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise
